"""Bounded file parsing with safe ZIP expansion and optional Arabic OCR."""
from __future__ import annotations

import hashlib
import os
from pathlib import Path, PurePosixPath
import shutil
import stat
import subprocess
import tempfile
import uuid
import zipfile

from .storage import now

MAX_UPLOAD = 100 * 1024 * 1024
MAX_EXPANDED = 250 * 1024 * 1024
MAX_MEMBERS = 200
MAX_TEXT = 2_000_000
SUPPORTED = {".pdf", ".docx", ".txt", ".csv", ".png", ".jpg", ".jpeg", ".tiff", ".tif", ".webp"}


def safe_name(name: str) -> str:
    name = name.replace("\\", "/").split("/")[-1]
    name = "".join(c for c in name if c >= " " and c not in "\x7f")
    return (name[:220] or "document")


def _ocr_image(path: Path) -> str:
    if not shutil.which("tesseract"):
        raise RuntimeError("محرك التعرف الضوئي غير مثبت")
    args = ["tesseract", str(path), "stdout", "-l", "ara+eng"]
    tessdata = os.environ.get("AUDIT_TESSDATA_DIR")
    if tessdata:
        args.extend(["--tessdata-dir", tessdata])
    process = subprocess.run(args, capture_output=True, timeout=45,
                             env={**os.environ, "OMP_THREAD_LIMIT": "1"})
    if process.returncode:
        raise RuntimeError("تعذر استخراج النص العربي؛ تحقق من ملفات لغة OCR")
    return process.stdout.decode("utf-8", errors="replace")[:MAX_TEXT]


def extract(path: Path, name: str) -> dict:
    ext = Path(name).suffix.lower()
    result = {"text": "", "page_count": None, "extraction_status": "unsupported", "warnings": []}
    try:
        if ext in {".txt", ".csv"}:
            data = path.read_bytes()
            for encoding in ["utf-8-sig", "utf-16", "cp1256"]:
                try:
                    result["text"] = data.decode(encoding)[:MAX_TEXT]
                    break
                except UnicodeError:
                    pass
            else:
                result["text"] = data.decode("utf-8", errors="replace")[:MAX_TEXT]
            result["extraction_status"] = "text"
            if len(data) > MAX_TEXT:
                result["extraction_status"] = "partial"
                result["warnings"].append("النص المستخرج محدود بمليوني حرف؛ راجع الملف الأصلي كاملاً")
        elif ext == ".docx":
            # A DOCX is itself a ZIP. Check expansion before handing it to the parser.
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED or len(archive.infolist()) > 3000:
                    raise ValueError("حجم محتوى Word بعد فك الضغط يتجاوز الحد المسموح")
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            joined = "\n".join(parts)
            result["text"] = joined[:MAX_TEXT]
            result["extraction_status"] = "partial" if len(joined) > MAX_TEXT else "text"
            if len(joined) > MAX_TEXT:
                result["warnings"].append("نص Word المستخرج جزئي؛ راجع الأصل كاملاً")
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path), strict=False)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("ملف PDF محمي بكلمة مرور")
            result["page_count"] = len(reader.pages)
            limit = min(150, len(reader.pages))
            pieces = []
            chars = 0
            scanned_gap = False
            for index in range(limit):
                source_page = reader.pages[index]
                page = source_page.extract_text() or ""
                pieces.append(f"\n[صفحة {index+1}]\n{page}")
                chars += len(page)
                resources = source_page.get("/Resources")
                if resources is not None:
                    resources = resources.get_object()
                    if resources.get("/XObject") and len(page.strip()) < 30:
                        scanned_gap = True
                if chars >= MAX_TEXT:
                    break
            result["text"] = "\n".join(pieces)[:MAX_TEXT] if chars else ""
            if chars >= 30:
                result["extraction_status"] = "partial" if limit < len(reader.pages) or chars >= MAX_TEXT or scanned_gap else "text"
                if result["extraction_status"] == "partial":
                    result["warnings"].append("النص المستخرج جزئي بسبب حدود الصفحات أو الحروف؛ راجع أصل الملف كاملاً")
            else:
                result["extraction_status"] = "needs_ocr"
                result["warnings"].append("الملف ممسوح ضوئياً؛ يلزم التحقق البشري من النص")
                # OCR is opt-in to keep uploads responsive. PDF/image originals remain usable evidence.
                if os.environ.get("AUDIT_OCR_ON_UPLOAD") == "1" and shutil.which("pdftoppm"):
                    page_limit = min(len(reader.pages), 5)
                    with tempfile.TemporaryDirectory(prefix="audit-ocr-") as temp:
                        prefix = str(Path(temp) / "page")
                        subprocess.run(["pdftoppm", "-r", "130", "-f", "1", "-l", str(page_limit),
                                        "-png", "-scale-to", "1800", str(path), prefix],
                                       capture_output=True, timeout=60, check=True)
                        ocr = []
                        for index, image in enumerate(sorted(Path(temp).glob("page-*.png")), 1):
                            ocr.append(f"\n[صفحة {index}]\n{_ocr_image(image)}")
                        result["text"] = "\n".join(ocr)[:MAX_TEXT]
                    result["extraction_status"] = "ocr_partial" if len(reader.pages) > page_limit else "ocr"
                    if len(reader.pages) > page_limit:
                        result["warnings"].append("OCR جزئي لأول خمس صفحات فقط؛ النص لا يثبت مراجعة كامل الملف")
        elif ext in SUPPORTED:
            from PIL import Image
            with Image.open(path) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("أبعاد الصورة كبيرة جداً؛ استخدم صورة لا تتجاوز 40 مليون بكسل")
                result["page_count"] = getattr(image, "n_frames", 1)
                image.verify()
            result["extraction_status"] = "needs_ocr"
            result["warnings"].append("صورة تحتاج قراءة بشرية أو استخراجاً ضوئياً")
            if os.environ.get("AUDIT_OCR_ON_UPLOAD") == "1":
                result["text"] = _ocr_image(path)
                result["extraction_status"] = "ocr"
    except Exception as exc:
        result["extraction_status"] = "failed"
        result["warnings"].append(str(exc)[:250] if isinstance(exc, (ValueError, RuntimeError))
                                  else "تعذر استخراج النص؛ أصل المستند محفوظ للمراجعة")
    return result


def create_document(path: Path, name: str, category: str, *, reference_year=None, title=None) -> dict:
    parsed = extract(path, name)
    with path.open("rb") as source:
        digest = hashlib.file_digest(source, "sha256").hexdigest()
    return {"id": path.stem, "name": safe_name(title or name), "original_name": safe_name(name),
            "category": category, "size": path.stat().st_size,
            "sha256": digest,
            "page_count": parsed["page_count"], "extraction_status": parsed["extraction_status"],
            "extracted_characters": len(parsed["text"]), "warnings": parsed["warnings"],
            "reference_year": reference_year, "created_at": now(),
            "_path": str(path), "_text": parsed["text"]}


def save_upload(stream, name: str, destination: Path, category: str) -> list[dict]:
    name = safe_name(name)
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED | {".zip"}:
        raise ValueError("نوع الملف غير مدعوم؛ استخدم PDF أو Word أو نصاً أو صورة أو ZIP")
    temp_path = destination / f"{uuid.uuid4().hex}{ext}"
    total = 0
    try:
        with temp_path.open("wb") as out:
            while block := stream.read(1024 * 1024):
                total += len(block)
                if total > MAX_UPLOAD:
                    raise ValueError("الحد الأقصى لكل ملف 100 ميغابايت")
                out.write(block)
        if not total:
            raise ValueError("الملف المرفوع فارغ")
        if ext != ".zip":
            return [create_document(temp_path, name, category)]
        return _expand_zip(temp_path, destination, category)
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        temp_path.unlink(missing_ok=True)
        raise ValueError("ملف ZIP غير صالح أو تالف") from exc
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    finally:
        if ext == ".zip":
            temp_path.unlink(missing_ok=True)


def _expand_zip(path: Path, destination: Path, category: str) -> list[dict]:
    documents, created = [], []
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_MEMBERS:
                raise ValueError("الأرشيف يتجاوز 200 ملف؛ قسّمه إلى أجزاء")
            if sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise ValueError("محتوى الأرشيف يتجاوز 250 ميغابايت بعد فك الضغط")
            for info in infos:
                posix = PurePosixPath(info.filename.replace("\\", "/"))
                if not posix.parts or posix.is_absolute() or ".." in posix.parts or ":" in posix.parts[0]:
                    raise ValueError("الأرشيف يحتوي مساراً غير آمن")
                if stat.S_ISLNK(info.external_attr >> 16):
                    raise ValueError("روابط الملفات داخل الأرشيف غير مسموحة")
                if info.flag_bits & 1:
                    raise ValueError("الأرشيف المشفر غير مدعوم")
                if info.file_size > MAX_UPLOAD or info.file_size / max(info.compress_size, 1) > 250:
                    raise ValueError("أحد ملفات الأرشيف يتجاوز حدود الحجم أو الضغط الآمن")
                if info.is_dir() or info.filename.startswith("__MACOSX/"):
                    continue
                if info.file_size == 0:
                    raise ValueError("الأرشيف يحتوي ملفاً فارغاً؛ أزل الملفات الفارغة ثم أعد الرفع")
                name = safe_name(info.filename)
                ext = Path(name).suffix.lower()
                if ext not in SUPPORTED:
                    raise ValueError(f"الأرشيف يحتوي نوعاً غير مدعوم: {name}")
                target = destination / f"{uuid.uuid4().hex}{ext}"
                created.append(target)
                written = 0
                with archive.open(info) as source, target.open("wb") as output:
                    while block := source.read(1024 * 1024):
                        written += len(block)
                        if written > MAX_UPLOAD or written > info.file_size:
                            raise ValueError("حجم ملف الأرشيف لا يطابق بياناته")
                        output.write(block)
                documents.append(create_document(target, name, category))
            if not documents:
                raise ValueError("لا يحتوي الأرشيف ملفات مدعومة")
        return documents
    except Exception:
        for target in created:
            target.unlink(missing_ok=True)
        raise
